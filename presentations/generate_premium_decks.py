from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

OUT = Path(__file__).resolve().parent / "premium"
OUT.mkdir(parents=True, exist_ok=True)

WHITE=RGBColor(255,255,255); TEXT=RGBColor(17,24,39); TEXT2=RGBColor(55,65,81)
MUTED=RGBColor(100,116,139); ACC=RGBColor(11,94,126); BLUE=RGBColor(37,99,235)
AMBER=RGBColor(245,158,11); VIOLET=RGBColor(124,58,237); LIGHT=RGBColor(247,250,252)
LIGHT2=RGBColor(241,245,249); MID=RGBColor(226,232,240); GREEN=RGBColor(22,163,74); RED=RGBColor(220,38,38)
COLORS=[ACC,BLUE,AMBER,VIOLET]

def wide(prs):
    prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)

def bg(slide):
    for x,y,w,h in [(10.9,-0.15,3,1.8),(-0.45,5.9,2.7,1.8)]:
        s=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid(); s.fill.fore_color.rgb=LIGHT2; s.line.color.rgb=LIGHT2

def footer(slide,n):
    ln=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.55), Inches(6.83), Inches(12.1), Inches(.015))
    ln.fill.solid(); ln.fill.fore_color.rgb=MID; ln.line.color.rgb=MID
    t=slide.shapes.add_textbox(Inches(.62), Inches(6.9), Inches(7), Inches(.2))
    p=t.text_frame.paragraphs[0]; p.text="ITS Umbria • VAPT 2026"; p.font.size=Pt(9.5); p.font.color.rgb=MUTED
    b=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(12), Inches(6.84), Inches(.55), Inches(.32))
    b.fill.solid(); b.fill.fore_color.rgb=LIGHT2; b.line.color.rgb=LIGHT2
    t=slide.shapes.add_textbox(Inches(12.08), Inches(6.89), Inches(.4), Inches(.18))
    p=t.text_frame.paragraphs[0]; p.text=str(n); p.font.size=Pt(10); p.font.bold=True; p.font.color.rgb=ACC; p.alignment=PP_ALIGN.CENTER

def chip(slide,x,y,w,text,color=ACC):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(.38))
    s.fill.solid(); s.fill.fore_color.rgb=LIGHT2; s.line.color.rgb=LIGHT2
    t=slide.shapes.add_textbox(Inches(x+.12), Inches(y+.08), Inches(w-.24), Inches(.18))
    p=t.text_frame.paragraphs[0]; p.text=text.upper(); p.font.size=Pt(11); p.font.bold=True; p.font.color.rgb=color

def header(slide,title,sub,day):
    bg(slide); chip(slide,.9,.62,1.15,day)
    bar=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.62), Inches(.58), Inches(.16), Inches(.82))
    bar.fill.solid(); bar.fill.fore_color.rgb=ACC; bar.line.color.rgb=ACC
    t=slide.shapes.add_textbox(Inches(.9), Inches(1.04), Inches(10.5), Inches(.85)); tf=t.text_frame
    p=tf.paragraphs[0]; p.text=title; p.font.size=Pt(24); p.font.bold=True; p.font.color.rgb=TEXT
    if sub:
        q=tf.add_paragraph(); q.text=sub; q.font.size=Pt(12); q.font.color.rgb=MUTED

def title_slide(prs,spec):
    s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); chip(s,.72,.82,1.28,spec["day"])
    t=s.shapes.add_textbox(Inches(.72), Inches(1.48), Inches(8.4), Inches(1.65)); tf=t.text_frame
    p=tf.paragraphs[0]; p.text=spec["title"]; p.font.size=Pt(31); p.font.bold=True; p.font.color.rgb=TEXT
    q=tf.add_paragraph(); q.text=spec["subtitle"]; q.font.size=Pt(17); q.font.color.rgb=TEXT2
    chip(s,.74,3.28,2.5,"moduli • "+spec["modules"],BLUE)
    t=s.shapes.add_textbox(Inches(.75), Inches(4.1), Inches(6), Inches(.9)); p=t.text_frame.paragraphs[0]
    p.text="Slide premium da proiezione: pochi testi, gerarchia chiara, alternanza naturale con il sito del corso."; p.font.size=Pt(13.5); p.font.color.rgb=MUTED
    for i,(label,col) in enumerate([("lezione",ACC),("laboratorio",BLUE),("debrief",AMBER)]):
        y=2+i*1.15
        b=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.1), Inches(y), Inches(3.1), Inches(.9)); b.fill.solid(); b.fill.fore_color.rgb=LIGHT; b.line.color.rgb=WHITE
        a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.1), Inches(y), Inches(.11), Inches(.9)); a.fill.solid(); a.fill.fore_color.rgb=col; a.line.color.rgb=col
        t=s.shapes.add_textbox(Inches(9.45), Inches(y+.28), Inches(2.4), Inches(.25)); p=t.text_frame.paragraphs[0]; p.text=label; p.font.size=Pt(18); p.font.bold=True; p.font.color.rgb=col
    footer(s,len(prs.slides))

def agenda(prs,spec):
    s=prs.slides.add_slide(prs.slide_layouts[6]); header(s,"Roadmap della giornata","Struttura del tempo e deliverable principali",spec["day"])
    xy=[(.78,1.8),(3.95,1.8),(.78,3.28),(3.95,3.28)]
    for i,((x,y),(a,b)) in enumerate(zip(xy,spec["agenda"])):
        box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.92), Inches(1.12)); box.fill.solid(); box.fill.fore_color.rgb=LIGHT; box.line.color.rgb=WHITE
        top=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(2.92), Inches(.08)); top.fill.solid(); top.fill.fore_color.rgb=COLORS[i]; top.line.color.rgb=COLORS[i]
        t=s.shapes.add_textbox(Inches(x+.2), Inches(y+.18), Inches(2.4), Inches(.75)); tf=t.text_frame
        p=tf.paragraphs[0]; p.text=a; p.font.size=Pt(17); p.font.bold=True; p.font.color.rgb=TEXT
        q=tf.add_paragraph(); q.text=b; q.font.size=Pt(11.5); q.font.color.rgb=MUTED
    box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.25), Inches(1.75), Inches(3.25), Inches(3.9)); box.fill.solid(); box.fill.fore_color.rgb=LIGHT; box.line.color.rgb=WHITE
    t=s.shapes.add_textbox(Inches(9.52), Inches(2), Inches(2.7), Inches(3.2)); tf=t.text_frame
    p=tf.paragraphs[0]; p.text="Output della giornata"; p.font.size=Pt(18); p.font.bold=True; p.font.color.rgb=ACC
    for item in spec["outputs"]:
        q=tf.add_paragraph(); q.text="• "+item; q.font.size=Pt(13); q.font.color.rgb=TEXT
    footer(s,len(prs.slides))

def cards(prs,spec,title,subtitle,items):
    s=prs.slides.add_slide(prs.slide_layouts[6]); header(s,title,subtitle,spec["day"])
    cols=3 if len(items)<=3 else 2; rows=(len(items)+cols-1)//cols; gap=.24; total=11.9; w=(total-gap*(cols-1))/cols; h=1.92 if rows==1 else 1.68
    for i,it in enumerate(items):
        r,c=divmod(i,cols); x=.75+c*(w+gap); y=1.9+r*(h+.32); col=COLORS[i%4]
        box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); box.fill.solid(); box.fill.fore_color.rgb=LIGHT; box.line.color.rgb=WHITE
        top=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(.08)); top.fill.solid(); top.fill.fore_color.rgb=col; top.line.color.rgb=col
        t=s.shapes.add_textbox(Inches(x+.18), Inches(y+.16), Inches(w-.3), Inches(h-.2)); tf=t.text_frame
        p=tf.paragraphs[0]; p.text=it[0]; p.font.size=Pt(17.5); p.font.bold=True; p.font.color.rgb=col
        for line in it[1:]:
            q=tf.add_paragraph(); q.text="• "+line; q.font.size=Pt(12.3); q.font.color.rgb=TEXT
    footer(s,len(prs.slides))

def process(prs,spec,title,subtitle,steps):
    s=prs.slides.add_slide(prs.slide_layouts[6]); header(s,title,subtitle,spec["day"])
    w=1.66; gap=.18; total=len(steps)*w+(len(steps)-1)*gap; x0=(13.333-total)/2; y=2.65
    for i,step in enumerate(steps):
        x=x0+i*(w+gap); box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1)); box.fill.solid(); box.fill.fore_color.rgb=LIGHT; box.line.color.rgb=WHITE
        t=s.shapes.add_textbox(Inches(x+.12), Inches(y+.34), Inches(w-.24), Inches(.35)); p=t.text_frame.paragraphs[0]; p.text=step; p.font.size=Pt(13.5); p.font.bold=True; p.font.color.rgb=TEXT; p.alignment=PP_ALIGN.CENTER
        if i<len(steps)-1:
            c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x+w), Inches(y+.5), Inches(x+w+gap), Inches(y+.5)); c.line.color.rgb=MID; c.line.width=Pt(2)
    footer(s,len(prs.slides))

def split(prs,spec,title,subtitle,left_title,left,right_title,right):
    s=prs.slides.add_slide(prs.slide_layouts[6]); header(s,title,subtitle,spec["day"])
    for x,tt,arr,col in [(.75,left_title,left,ACC),(6.58,right_title,right,BLUE)]:
        b=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.95), Inches(5), Inches(4.45)); b.fill.solid(); b.fill.fore_color.rgb=LIGHT; b.line.color.rgb=WHITE
        t=s.shapes.add_textbox(Inches(x+.22), Inches(2.18), Inches(4.52), Inches(3.95)); tf=t.text_frame
        p=tf.paragraphs[0]; p.text=tt; p.font.size=Pt(18); p.font.bold=True; p.font.color.rgb=col
        for item in arr:
            q=tf.add_paragraph(); q.text="• "+item; q.font.size=Pt(12.8); q.font.color.rgb=TEXT
    footer(s,len(prs.slides))

def checklist(prs,spec,title,subtitle,items):
    s=prs.slides.add_slide(prs.slide_layouts[6]); header(s,title,subtitle,spec["day"])
    b=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.86), Inches(1.95), Inches(11.55), Inches(4.45)); b.fill.solid(); b.fill.fore_color.rgb=LIGHT; b.line.color.rgb=WHITE
    t=s.shapes.add_textbox(Inches(1.18), Inches(2.25), Inches(10.8), Inches(3.8)); tf=t.text_frame
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text="☐ "+item; p.font.size=Pt(15.5); p.font.color.rgb=TEXT
    footer(s,len(prs.slides))

def quote(prs,spec,text,small):
    s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); chip(s,.72,.72,1.15,spec["day"])
    bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.08), Inches(1.92), Inches(.18), Inches(2.4)); bar.fill.solid(); bar.fill.fore_color.rgb=ACC; bar.line.color.rgb=ACC
    t=s.shapes.add_textbox(Inches(1.5), Inches(1.88), Inches(9.8), Inches(2.8)); tf=t.text_frame
    p=tf.paragraphs[0]; p.text=text; p.font.size=Pt(27); p.font.bold=True; p.font.color.rgb=TEXT
    q=tf.add_paragraph(); q.text=small; q.font.size=Pt(14); q.font.color.rgb=MUTED
    footer(s,len(prs.slides))

def site(prs,spec):
    s=prs.slides.add_slide(prs.slide_layouts[6]); header(s,"Quando passare al sito del corso","Le slide danno ritmo. Il sito contiene i dettagli operativi.",spec["day"])
    data=[("qui in slide",["ritmo della giornata","concetti chiave","mappe e diagrammi","output attesi"]),("sul sito",["lezione completa","dossier di laboratorio","artefatti e request","quiz e PDF"]),("durante il lab",["consultare il sito","annotare nel taccuino","chiedere evidenze, non solo risposte","usare i deliverable come guida"])]
    for i,(tt,arr) in enumerate(data):
        x=.8+i*4.08; b=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2), Inches(3.45), Inches(3.85)); b.fill.solid(); b.fill.fore_color.rgb=LIGHT; b.line.color.rgb=WHITE
        t=s.shapes.add_textbox(Inches(x+.18), Inches(2.25), Inches(3), Inches(3.2)); tf=t.text_frame
        p=tf.paragraphs[0]; p.text=tt.upper(); p.font.size=Pt(16); p.font.bold=True; p.font.color.rgb=COLORS[i]
        for item in arr:
            q=tf.add_paragraph(); q.text="• "+item; q.font.size=Pt(12.6); q.font.color.rgb=TEXT
    footer(s,len(prs.slides))

DECKS=[
{"day":"DAY 01","file":"day-01-mindset-scope-premium.pptx","modules":"00 + 01","title":"Giornata 1 — Mindset, incarico e scope","subtitle":"Facts, hypotheses, scope e Rules of Engagement","agenda":[("baseline","livello iniziale e linguaggio comune"),("mindset","facts, hypotheses, evidence"),("engagement","cosa significa fare un VAPT"),("client kickoff","scope, limiti e deliverable")],"outputs":["Facts / Hypotheses / Questions","Scope & RoE","Kickoff notes"],"slides":[
("quote","Il corso non insegna a lanciare tool. Insegna a prendere decisioni tecniche supportate da evidenze.","La frase-guida della prima giornata."),
("cards","Tre parole chiave","Separare ciò che so da ciò che sto solo ipotizzando",[("fact","osservazione supportata da evidenza"),("hypothesis","spiegazione plausibile ma non ancora provata"),("evidence","ciò che sostiene una conclusione")]),
("process","Come ragiona un pentester","Metodo generale del corso",["osservo","ipotesi","test","evidenza","conclusione"]),
("cards","VA vs Penetration Test","Due attività diverse, stesso rigore",[("VA","identifica e prioritizza debolezze","non sempre richiede exploitation"),("PT","verifica impatto","può includere PoC"),("obiettivo comune","capire il rischio","aiutare a correggere")]),
("split","Kickoff con il cliente","Il lavoro professionale parte prima dei test","domande da fare",["perimetro autorizzato","domini / host / ambienti","account e ruoli","out of scope","contatti ed escalation"],"cose da chiarire",["exploitation consentita?","DoS / distruttivi vietati?","dati reali o test?","deliverable attesi?","stop condition?"]),
("process","Client Kickoff","Trasformare una richiesta vaga in un incarico gestibile",["brief","domande","scope","RoE","deliverable"]),
("check","Cosa devono portarsi via","Deliverable della giornata",["distinguere facts, hypotheses ed evidence","spiegare VA vs PT","fare domande prima di toccare un target","scrivere scope e RoE minime"]),("site",)]},
{"day":"DAY 02","file":"day-02-foundations-recon-premium.pptx","modules":"02 + 03","title":"Giornata 2 — Fondamentali tecnici + Recon","subtitle":"DNS, HTTP, sessione e attack surface","agenda":[("fondamentali","DNS → TCP → TLS → HTTP"),("request anatomy","capire davvero una richiesta"),("recon","host, porte, servizi e hostname"),("attack surface","mappa utile, non output grezzo")],"outputs":["Request Anatomy","Attack Surface Inventory"],"slides":[
("process","Dal browser all’applicazione","La request attraversa livelli diversi",["DNS","IP/TCP","TLS","HTTP","app"]),
("cards","Anatomia minima di una request","Cosa leggere nel traffico reale",[("linea iniziale","metodo","path","versione HTTP"),("header","Host","Cookie","Content-Type"),("body","form","JSON","file upload")]),
("split","Che cosa voglio sapere da una request?","Una request genera già domande di sicurezza","osservare",["chi parla con chi","oggetto richiesto","input controllabili","ruolo","response"],"domandarsi",["cosa identifica l’utente?","cosa identifica la risorsa?","cosa cambia senza sessione?","quale parte può cambiare?","questa info cambia il piano?"]),
("cards","Recon = mappatura","Non è una gara di scanner",[("asset","hostname","IP","relazioni"),("esposizione","porte","servizi","tecnologie dichiarate"),("priorità","cosa approfondire","cosa è fuori scope")]),
("process","Build the Attack Surface","Dall’output alla mappa",["hostname","porte","servizi","contenuto","ipotesi","next test"]),
("split","Strumenti: quale domanda rispondono?","I tool sono strumenti, non conclusioni","esempi",["dig → risoluzione hostname","nmap → servizi esposti","curl → risposta HTTP","robots.txt → aree dichiarate"],"errori",["comando prima della domanda","banner = verità assoluta","porta aperta = vulnerabilità","testare tutto senza priorità"]),
("check","Output della giornata","Cosa ci aspettiamo nei taccuini",["leggere una request end-to-end","distinguere host / hostname / servizio / app","creare Attack Surface Inventory","segnalare correttamente asset fuori scope"]),("site",)]},
{"day":"DAY 03","file":"day-03-va-web-mapping-premium.pptx","modules":"04 + 05","title":"Giornata 3 — Vulnerability Assessment + Web Mapping","subtitle":"Scanner, validazione, proxy e application mapping","agenda":[("VA","scanner ≠ assessment"),("validazione","CVE, prerequisiti, contesto"),("proxy","browser ↔ proxy ↔ web/API"),("mapping","funzioni, ruoli, input, oggetti")],"outputs":["Validation table","Application Attack Surface Map"],"slides":[
("process","Da alert a finding","Un assessment aggiunge ragionamento umano",["scanner","evidenza","prerequisiti","verifica","contesto","conclusione"]),
("cards","Tre sigle da usare bene","Servono solo se aiutano a spiegare il finding",[("CVE","vulnerabilità specifica nota","non prova da sola il problema"),("CWE","classe di debolezza","aiuta a capire la causa"),("CVSS","severity tecnica","non sostituisce il rischio business")]),
("cards","False positive vs false negative","Scanner = input, non fonte di verità",[("false positive","il tool segnala","problema non applicabile"),("false negative","problema esiste","tool non lo vede"),("lezione","validare","contestualizzare")]),
("process","Il proxy rende osservabile l’app","Una variabile per volta",["browser","proxy","request","response","replay"]),
("cards","Che cosa mappare","Prima si mappa, poi si testa",[("funzione","cosa fa la UI","crea / legge / modifica"),("richiesta","metodo","endpoint","sessione"),("sicurezza","oggetto","ruolo","next test")]),
("split","Checklist di web mapping","Base dei moduli 06 e 07","osservare",["endpoint e method","input controllabili","cookie / token","oggetti","transizioni di stato"],"chiedersi",["UI è unico accesso?","quale ID identifica la risorsa?","cosa cambia con un input?","quale ruolo?","finding o solo mappa?"]),
("check","Deliverable attesi","Ponte verso testing applicativo",["validation table","application attack surface map","next test motivati"]),("site",)]},
{"day":"DAY 04","file":"day-04-input-handling-injection-premium.pptx","modules":"06","title":"Giornata 4 — Input Handling & Injection","subtitle":"Dal parametro alla prova minima","agenda":[("input","dati non fidati e trust boundary"),("contesto","SQL, HTML, shell, path"),("metodo","baseline → test → differenza"),("lab","Input to Evidence")],"outputs":["Input evidence sheet","Finding con remediation"],"slides":[
("quote","La domanda iniziale non è ‘quale payload provo?’. È: ‘dove finisce questo input e chi lo interpreta?’","Se il contesto è sbagliato, anche il test lo sarà."),
("cards","Input non fidato","Capire prima di testare",[("entry point","query string","path parameter","form / JSON"),("altri canali","header","cookie","nome file"),("rischio","trust boundary","significato inatteso")]),
("cards","Famiglie di problemi","Comprendere il meccanismo",[("SQLi","input entra nella query","serve evidenza coerente"),("XSS","input finisce in HTML/JS","browser interpreta codice"),("command / path","input altera comandi o percorsi","PoC minima")]),
("process","Metodo operativo","Una variabile per volta",["input","baseline","test minimo","differenza","conferma","impatto"]),
("cards","Detection vs impact","Fermarsi quando l’evidenza è sufficiente",[("detection","cerco un segnale","differenza riproducibile"),("impact","dimostro conseguenza","non sempre necessario"),("regola","minimo impatto","stop condition")]),
("split","Errori comuni","Il corso premia metodo","da evitare",["liste di payload","500 = prova","cambiare 5 cose insieme","remediation generiche"],"da fare",["baseline chiara","differenze coerenti","ripetere per conferma","remediation sulla causa"]),
("check","Deliverable della giornata","Cosa vogliamo vedere",["input e contesto","baseline + test minimo","differenza + conferma","impatto supportato","remediation causale"]),("site",)]},
{"day":"DAY 05","file":"day-05-auth-api-security-premium.pptx","modules":"07","title":"Giornata 5 — Authentication, Authorization & API Security","subtitle":"Two Users, One Object","agenda":[("authn","chi sei?"),("sessione","come il server ti riconosce?"),("authz","cosa puoi fare e su quale oggetto?"),("API testing","confronti tra utenti e ruoli")],"outputs":["API map","Authorization test evidence","Finding access-control/API"],"slides":[
("cards","Tre concetti distinti","Non confonderli",[("authentication","verifica identità","login"),("session management","lega richieste e utente","cookie / token"),("authorization","azione + oggetto + ruolo","cuore della lezione")]),
("cards","Access control","Il controllo deve essere server-side",[("orizzontale","Alice accede a Bob","stesso ruolo"),("verticale","utente standard usa admin","ruolo diverso"),("function level","UI non basta","server deve rifiutare")]),
("process","Two Users, One Object","Test cross-user",["Alice","ordine 1001","Bob","ordine 1002","cross-user"]),
("split","Cosa mappare in una API","I confronti generano i test","struttura",["metodo e path","sessione / bearer","object ID","effetto","ruolo"],"confronti",["anonimo vs auth","Alice vs Bob","user vs admin","GET vs PUT/DELETE","UI vs request diretta"]),
("cards","JWT: idea operativa","Capire il flusso prima dei trucchi",[("header.payload.signature","header e payload leggibili","non è cifratura"),("domande utili","quali claim?","cosa usa il server?"),("errore","token hacking senza ipotesi","prima capire authz")]),
("check","Deliverable della giornata","Focus sulla decisione server-side",["mappa di endpoint","test cross-user / cross-role","identità vs sessione vs risorsa","PoC minima"]),("site",)]},
{"day":"DAY 06","file":"day-06-exploitation-postex-reporting-ai-premium.pptx","modules":"08 + 09 + 10 + 11","title":"Giornata 6 — Exploitation, Post-Exploitation, Reporting e AI","subtitle":"Decisioni, non spettacolo","agenda":[("08","quando una PoC aggiunge valore"),("09","security context e privilege boundary"),("10","da evidence a finding"),("11","AI propone, umano verifica")],"outputs":["Exploitation decision sheet","Attack path","Finding professionale","AI claim review"],"slides":[
("process","Decision gate prima di exploitare","Se manca un punto critico, non si parte",["finding confermato?","vale la pena?","RoE ok?","stop condition?","evidenza sufficiente?"]),
("cards","Vulnerability / exploit / payload","Non sono sinonimi",[("vulnerability","debolezza nel sistema"),("exploit","tecnica che la sfrutta"),("payload","effetto successivo")]),
("cards","Post-exploitation: domande giuste","Misurare propagazione e impatto",[("chi sono?","utente","host","privilegi"),("quale confine?","privilege boundary","configurazione"),("quando fermarmi?","rischio dimostrato","cleanup")]),
("cards","Un finding professionale","Il report è il prodotto finale",[("causa","non solo sintomo"),("evidenza","riproducibile","minimizzata"),("remediation","agisce sulla causa")]),
("cards","AI nel corso","AI output ≠ evidence",[("utile per","riassumere note","checklist"),("non sostituisce","evidenza","fonte primaria"),("regola","verifica umana","scope invariato")]),
("split","Briefing del capstone","Preparare la giornata successiva","gli studenti dovranno",["scope","attack surface","validare almeno un problema","stop condition","finding + summary"],"noi osserviamo",["metodo","prioritizzazione","note","uso hint","chiarezza report"]),("site",)]},
{"day":"DAY 07","file":"day-07-capstone-premium.pptx","modules":"12","title":"Giornata 7 — Capstone UmbraMarket","subtitle":"Mini-assessment end-to-end","agenda":[("briefing","scope, obiettivi e deliverable"),("planning","attack surface e priorità"),("testing","validation ed evidence"),("reporting","finding + executive summary")],"outputs":["Attack Surface","Structured Notes","Finding(s)","Evidence Pack"],"slides":[
("quote","Il capstone non misura quanti tool sapete usare. Misura se sapete condurre un assessment professionale.","Tutto ciò che serve è già comparso nei moduli precedenti."),
("process","Flusso consigliato","Il processo conta più della velocità",["scope","surface","ipotesi","test","evidenza","finding","summary"]),
("cards","Che cosa consegnare","Artefatti leggibili",[("notes","facts","hypotheses","tests"),("finding","evidence","impatto","remediation"),("summary","rischi","priorità","linguaggio cliente")]),
("cards","Regola sugli hint","Metodo prima della soluzione",[("level 0","nessun hint"),("level 1–2","domanda sul metodo","rileggi evidence"),("level 3–4","area / concetto","passaggio guidato")]),
("split","Priorità nel capstone","Meglio un buon finding che cinque alert","proteggere",["metodo e note","finding validato","stop condition","report comprensibile"],"evitare",["fare tutto a metà","continuare senza valore","scanner = finding","inventare impatti"]),
("check","Checklist finale","Prima di chiudere",["scope chiaro","fatto vs ipotesi","test motivato","evidenza minima","passare al reporting"]),("site",)]},
{"day":"DAY 08","file":"day-08-closing-report-presentation-premium.pptx","modules":"12 + wrap-up","title":"Giornata 8 — Chiusura, report e client presentation","subtitle":"Peer review, executive summary e debrief finale","agenda":[("chiusura finding","pulire e completare evidenze"),("peer review","controllo incrociato"),("client presentation","rischio e priorità"),("debrief","cosa resta alla classe")],"outputs":["Mini report finale","Executive summary","Presentazione breve"],"slides":[
("cards","Executive summary vs technical finding","Due pubblici, stesso rigore",[("executive summary","cosa testato","rischi","priorità"),("finding tecnico","causa","evidenza","remediation"),("punto chiave","pubblici diversi","stessa verità")]),
("process","Da note grezze a report utile","Ordinare il lavoro è professionalità",["notes","facts","finding","summary","presentation"]),
("check","Peer review checklist","Controllo incrociato",["titolo descrive il problema","descrizione spiega la causa","evidenza sufficiente","riproduzione chiara","impatto supportato","remediation causale"]),
("split","Mini presentazione al cliente","Farsi capire","in 2 minuti",["perimetro","rischio principale","evidenza high-level","priorità","limiti"],"evitare",["leggere slide","dettagli inutili","tool protagonista","sovrastimare il rischio"]),
("cards","Cosa saper fare a fine corso","Competenze visibili",[("partire","scope","regole","surface"),("testare","ipotesi","evidence","stop condition"),("comunicare","finding","summary","priorità")]),
("quote","Mi danno un target autorizzato che non conosco. So come partire, cosa osservare, formulare ipotesi, verificarle e spiegare il rischio.","La frase finale del percorso."),("site",)]}
]

for spec in DECKS:
    prs=Presentation(); wide(prs); title_slide(prs,spec); agenda(prs,spec)
    for sl in spec["slides"]:
        kind=sl[0]
        if kind=="quote": quote(prs,spec,sl[1],sl[2])
        elif kind=="cards": cards(prs,spec,sl[1],sl[2],sl[3])
        elif kind=="process": process(prs,spec,sl[1],sl[2],sl[3])
        elif kind=="split": split(prs,spec,sl[1],sl[2],sl[3],sl[4],sl[5],sl[6])
        elif kind=="check": checklist(prs,spec,sl[1],sl[2],sl[3])
        elif kind=="site": site(prs,spec)
    prs.save(OUT/spec["file"])
    print("generated", OUT/spec["file"])
